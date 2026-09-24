from io import BytesIO
from typing import TYPE_CHECKING

from datalab_commons.files.media import (
    DOCX_MEDIA_TYPE,
    PPTX_MEDIA_TYPE,
    XLSX_MEDIA_TYPE,
    is_text_media_type,
)

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph

MAX_EXTRACTED_CHARS = 200_000
MAX_SHEET_ROWS = 1_000
TRUNCATION_NOTICE = "\n\n[conteúdo truncado]"

EXTRACTABLE_MEDIA_TYPES = frozenset({DOCX_MEDIA_TYPE, XLSX_MEDIA_TYPE, PPTX_MEDIA_TYPE})


class UnsupportedFileType(ValueError):
    def __init__(self, media_type: str) -> None:
        self.media_type = media_type
        super().__init__(f"Cannot extract text from {media_type}")


def decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def to_utf8(content: bytes) -> bytes:
    try:
        content.decode("utf-8")
    except UnicodeDecodeError:
        return decode_text(content).encode("utf-8")
    return content


def extract_text(content: bytes, media_type: str, *, max_chars: int = MAX_EXTRACTED_CHARS) -> str:
    if is_text_media_type(media_type):
        return _truncate(decode_text(content), max_chars)
    if media_type == DOCX_MEDIA_TYPE:
        return _truncate(_from_docx(content), max_chars)
    if media_type == XLSX_MEDIA_TYPE:
        return _truncate(_from_xlsx(content), max_chars)
    if media_type == PPTX_MEDIA_TYPE:
        return _truncate(_from_pptx(content), max_chars)
    raise UnsupportedFileType(media_type)


def _truncate(text: str, max_chars: int) -> str:
    return text if len(text) <= max_chars else text[:max_chars] + TRUNCATION_NOTICE


def _from_docx(content: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(content))
    blocks: list[str] = []
    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            blocks.append(_docx_paragraph(Paragraph(element, document)))
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            blocks.append(_markdown_table([[cell.text.strip() for cell in row.cells] for row in table.rows]))
    return "\n\n".join(block for block in blocks if block)


def _docx_paragraph(paragraph: "Paragraph") -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style = (getattr(paragraph.style, "name", "") or "").lower()
    if style.startswith("heading"):
        level = "".join(char for char in style if char.isdigit()) or "1"
        return f"{'#' * min(int(level), 6)} {text}"
    if style.startswith("list"):
        return f"- {text}"
    return text


def _from_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for worksheet in workbook.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in worksheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True)
            ]
            rows = [row for row in rows if any(cell for cell in row)]
            if not rows:
                continue
            sheets.append(f"## {worksheet.title}\n\n{_markdown_table(rows)}")
        return "\n\n".join(sheets)
    finally:
        workbook.close()


def _from_pptx(content: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(BytesIO(content))
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if shape.has_table:
                lines.append(_markdown_table([[cell.text.strip() for cell in row.cells] for row in shape.table.rows]))
        if lines:
            slides.append(f"## Slide {number}\n\n" + "\n\n".join(lines))
    return "\n\n".join(slides)


def _markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    header, *body = rows
    width = max(len(row) for row in rows)
    header = header + [""] * (width - len(header))
    lines = [_markdown_row(header), _markdown_row(["---"] * width)]
    lines.extend(_markdown_row(row + [""] * (width - len(row))) for row in body)
    return "\n".join(lines)


def _markdown_row(cells: list[str]) -> str:
    return "| " + " | ".join(cell.replace("|", "\\|").replace("\n", " ") for cell in cells) + " |"
