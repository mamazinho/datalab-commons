from io import BytesIO

import pytest

from datalab_commons.files import (
    DOCX_MEDIA_TYPE,
    PPTX_MEDIA_TYPE,
    XLSX_MEDIA_TYPE,
    UnsupportedFileType,
    decode_text,
    extract_text,
    to_utf8,
)


def build_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Relatório de campanha", level=1)
    document.add_paragraph("A campanha de verão cresceu 18%.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Canal"
    table.cell(0, 1).text = "CTR"
    table.cell(1, 0).text = "Meta"
    table.cell(1, 1).text = "2%"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def build_xlsx(rows: int = 3) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendas"
    sheet.append(["Canal", "Receita"])
    for number in range(rows):
        sheet.append([f"Canal {number}", 100 + number])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Plano de mídia"
    slide.placeholders[1].text = "Investimento de R$ 20 mil"
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


class TestDecodeText:
    @pytest.mark.parametrize(
        ("content", "expected"),
        [
            pytest.param("Ação".encode(), "Ação", id="utf-8"),
            pytest.param("﻿Ação".encode(), "Ação", id="utf-8-with-excel-bom"),
            pytest.param("Ação".encode("cp1252"), "Ação", id="windows-1252-from-excel"),
        ],
    )
    def test_decodes_whatever_encoding_the_file_came_in(self, content, expected):
        assert decode_text(content).lstrip("﻿") == expected

    def test_to_utf8_leaves_a_utf8_file_byte_for_byte(self):
        content = "Ação,1".encode()

        assert to_utf8(content) is content

    def test_to_utf8_converts_what_the_provider_could_not_decode(self):
        assert to_utf8("Ação,1".encode("cp1252")) == "Ação,1".encode()


class TestExtractText:
    def test_reads_a_docx_keeping_headings_and_tables(self):
        text = extract_text(build_docx(), DOCX_MEDIA_TYPE)

        assert "# Relatório de campanha" in text
        assert "A campanha de verão cresceu 18%." in text
        assert "| Canal | CTR |" in text

    def test_reads_every_sheet_of_a_xlsx_as_a_table(self):
        text = extract_text(build_xlsx(), XLSX_MEDIA_TYPE)

        assert "## Vendas" in text
        assert "| Canal | Receita |" in text
        assert "| Canal 0 | 100 |" in text

    def test_reads_a_pptx_slide_by_slide(self):
        text = extract_text(build_pptx(), PPTX_MEDIA_TYPE)

        assert "## Slide 1" in text
        assert "Plano de mídia" in text

    def test_reads_a_text_file_in_any_encoding(self):
        assert "Ação" in extract_text("Ação,1".encode("cp1252"), "text/csv")

    def test_truncates_so_one_attachment_cannot_eat_the_whole_context(self):
        text = extract_text(build_xlsx(rows=200), XLSX_MEDIA_TYPE, max_chars=200)

        assert len(text) < 300
        assert text.endswith("[conteúdo truncado]")

    def test_refuses_a_format_it_cannot_read(self):
        with pytest.raises(UnsupportedFileType):
            extract_text(b"PK\x03\x04", "application/zip")
